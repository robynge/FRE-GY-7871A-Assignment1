"""Create the Word companion from the same Markdown used by the main report."""
import json
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs'
cutoff = json.loads((OUT / 'audit.json').read_text())['sample_end']
path = OUT / f'ARK_Holdings_Sentiment_and_Uncertainty_Report_{cutoff}.docx'
pandoc = shutil.which('pandoc')
if not pandoc:
    raise RuntimeError('Pandoc is required for native Word equations.')
subprocess.run([pandoc, str(ROOT / 'REPORT.md'), '--from=markdown+tex_math_dollars',
                '--to=docx', '--standalone', '--resource-path=' + str(ROOT),
                '--output=' + str(path)], check=True)
doc = Document(path)
styles = {style.name: style for style in doc.styles}
sec = doc.sections[0]
sec.page_width, sec.page_height = Pt(595.276), Pt(841.89)
sec.top_margin, sec.bottom_margin = Pt(34), Pt(35)
sec.left_margin, sec.right_margin = Pt(40), Pt(40)
sec.header_distance, sec.footer_distance = Pt(13), Pt(14)
for style in doc.styles:
    if style.type == 1:
        style.font.name = 'Arial'
        style.font.size = Pt(9)
        style.font.color.rgb = RGBColor.from_string('0A0A23')
        style.paragraph_format.line_spacing = Pt(12)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        style.paragraph_format.space_after = Pt(6)
        for border in style.element.findall('.//' + qn('w:pBdr')):
            border.getparent().remove(border)
for name in ['Heading 1', 'Heading 2', 'Heading 3']:
    styles[name].font.size = Pt(11)
    styles[name].font.bold = True
    styles[name].paragraph_format.keep_with_next = True
    styles[name].paragraph_format.space_before = Pt(12)
styles['Title'].font.size = Pt(19)
styles['Title'].font.bold = True
styles['Title'].paragraph_format.line_spacing = Pt(23)
styles['Title'].paragraph_format.space_after = Pt(9)
doc.paragraphs[0].style = styles['Title']
doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
for p in doc.paragraphs:
    if p.text == 'Sources':
        p.paragraph_format.page_break_before = True
    if p.text.startswith('Panel B') or p.text.endswith(':'):
        p.paragraph_format.keep_with_next = True
    if p._p.findall('.//' + qn('w:drawing')):
        p.paragraph_format.keep_with_next = True
    if p.style.name in ['Caption', 'Image Caption']:
        p.style.font.size = Pt(8)
        p.style.font.color.rgb = RGBColor.from_string('676777')
    for run in p.runs:
        run.font.name = 'Arial'
for rel in doc.part.rels.values():
    if rel.reltype.endswith('/hyperlink'):
        for h in doc.element.findall('.//' + qn('w:hyperlink')):
            if h.get(qn('r:id')) == rel.rId:
                for run in h.findall(qn('w:r')):
                    props = run.find(qn('w:rPr'))
                    if props is None:
                        props = OxmlElement('w:rPr'); run.insert(0, props)
                    color = OxmlElement('w:color'); color.set(qn('w:val'), '6542CE'); props.append(color)

widths = [[63,113,113,113,113], [286,57,62,110], [299,64,76,76], [299,64,76,76], [299,64,76,76],
          [38,137,43,59,59,59,60,60], [35,163,65,182,70],
          [40,75,66,52,52,77,65,88], [135,190,190], [45,52,52,116,70,70,110],
          [40,50,105,60,60,70,65,65]]
assert len(doc.tables) == len(widths), 'Unexpected report table count'
for t, sizes in zip(doc.tables, widths):
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    sizes = [v * 515 / sum(sizes) for v in sizes]
    for col, width in zip(t.columns, sizes):
        col.width = Pt(width)
    for i, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            pr.append(OxmlElement('w:tblHeader'))
        for j, cell in enumerate(row.cells):
            cell.width = Pt(sizes[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'F0ECFF' if i == 0 else ('FAF9FE' if i % 2 else 'FFFFFF')); cp.append(shade)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el = OxmlElement('w:' + edge); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'DDD6F0'); borders.append(el)
            cp.append(borders)
            margins = OxmlElement('w:tcMar')
            for side, size in [('top',60),('bottom',60),('left',80),('right',80)]:
                el = OxmlElement('w:' + side); el.set(qn('w:w'), str(size)); el.set(qn('w:type'), 'dxa'); margins.append(el)
            cp.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(9.5)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
                p.paragraph_format.keep_with_next = i == 0 or (len(t.rows) <= 15 and i < len(t.rows) - 1)
                numeric = bool(re.match(r'^[+−-]?[\d.,]+$', cell.text)) or bool(p._p.findall('.//' + qn('m:oMath')))
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i and numeric else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = 'Arial'; run.font.size = Pt(7.5); run.font.bold = i == 0
                for run in p._p.findall('.//' + qn('m:r')):
                    pr = run.find(qn('w:rPr'))
                    if pr is None:
                        pr = OxmlElement('w:rPr'); run.insert(0, pr)
                    size = OxmlElement('w:sz'); size.set(qn('w:val'), '15'); pr.append(size)
    before = t._tbl.getprevious()
    if before is not None and before.tag == qn('w:p'):
        pr = before.find(qn('w:pPr'))
        if pr is None:
            pr = OxmlElement('w:pPr'); before.insert(0, pr)
        pr.append(OxmlElement('w:keepNext'))
for shape in doc.inline_shapes:
    ratio = int(shape.height) / int(shape.width)
    shape.width = Pt(515)
    shape.height = Pt(515 * ratio)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run(f'ARK holdings   |   {cutoff}   |   ')
run.font.size = Pt(7)
run.font.color.rgb = RGBColor.from_string('676777')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
doc.core_properties.title = 'Uncertainty and sentiment in financial reports'
doc.core_properties.author = 'Research'
assert len(doc.element.findall('.//' + qn('m:oMath'))) >= 12, 'Native equations were not preserved'
doc.save(path)
print(path)
